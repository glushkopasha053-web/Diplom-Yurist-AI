
from pathlib import Path
from pypdf import PdfReader
from docx import Document


MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 МБ
MAX_TEXT_LENGTH = 100_000         # максимум 100 000 символів


def extract_document_text(file_storage):
    """
    Безпечно перевіряє формат і розмір завантаженого файла
    та витягує текст із PDF або DOCX.
    """

    if not file_storage or not file_storage.filename:
        raise ValueError("Не вибрано файл.")

    filename = Path(file_storage.filename).name
    extension = Path(filename).suffix.lower()

    if extension not in (".pdf", ".docx"):
        raise ValueError("Дозволені лише файли PDF і DOCX.")

    # Читаємо файл у пам'ять та обмежуємо розмір
    file_storage.stream.seek(0)
    file_bytes = file_storage.stream.read(MAX_FILE_SIZE + 1)

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError("Файл завеликий. Максимальний розмір — 10 МБ.")

    if not file_bytes:
        raise ValueError("Файл порожній.")

    from io import BytesIO
    stream = BytesIO(file_bytes)

    try:
        if extension == ".pdf":
            reader = PdfReader(stream)

            if reader.is_encrypted:
                raise ValueError(
                    "PDF захищений паролем. Завантажте незахищений файл."
                )

            parts = []

            for page in reader.pages:
                parts.append(page.extract_text() or "")

            text = "\n".join(parts)

        else:
            document = Document(stream)

            parts = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            # Також читаємо текст із таблиць Word
            for table in document.tables:
                for row in table.rows:
                    parts.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

            text = "\n".join(parts)

    except ValueError:
        raise
    except Exception as error:
        raise ValueError(
            "Не вдалося прочитати файл. Перевірте, чи він не пошкоджений."
        ) from error

    text = text.strip()

    if not text:
        raise ValueError(
            "Не знайдено текст. Можливо, PDF є сканом без розпізнаного тексту."
        )

    return text[:MAX_TEXT_LENGTH]
